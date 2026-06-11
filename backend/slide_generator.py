import io
import os
import requests
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

def download_logo(platform):
    logos = {
        "Facebook": "https://upload.wikimedia.org/wikipedia/commons/thumb/0/05/Facebook_Logo_%282019%29.png/1024px-Facebook_Logo_%282019%29.png",
        "Instagram": "https://upload.wikimedia.org/wikipedia/commons/thumb/e/e7/Instagram_logo_2016.svg/2048px-Instagram_logo_2016.svg.png",
        "TikTok": "https://upload.wikimedia.org/wikipedia/en/thumb/a/a9/TikTok_logo.svg/1024px-TikTok_logo.svg.png",
        "YouTube": "https://upload.wikimedia.org/wikipedia/commons/thumb/0/09/YouTube_full-color_icon_%282017%29.svg/1024px-YouTube_full-color_icon_%282017%29.svg.png",
        "LinkedIn": "https://upload.wikimedia.org/wikipedia/commons/thumb/c/ca/LinkedIn_logo_initials.png/1024px-LinkedIn_logo_initials.png"
    }
    url = logos.get(platform)
    if not url:
        return None
    
    os.makedirs("backend/assets", exist_ok=True)
    filepath = f"backend/assets/{platform}.png"
    if not os.path.exists(filepath):
        try:
            r = requests.get(url, stream=True)
            if r.status_code == 200:
                with open(filepath, 'wb') as f:
                    for chunk in r: f.write(chunk)
        except Exception as e:
            print("Logo download failed:", e)
            return None
    return filepath

def _set_cell_text(cell, text, bold=False, color=None, size=Pt(11), align=PP_ALIGN.CENTER):
    text_frame = cell.text_frame
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = 'Open Sans'
    p.font.bold = bold
    p.font.size = size
    if color:
        p.font.color.rgb = color
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE

def _add_subtext(cell, text, color=RGBColor(100, 100, 100), size=Pt(9)):
    p = cell.text_frame.add_paragraph()
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = 'Open Sans'
    p.font.color.rgb = color
    p.font.size = size

def generate_platform_highlights_slide(platforms_data, prev_month_label):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    slide_layout = prs.slide_layouts[6] 
    slide = prs.slides.add_slide(slide_layout)
    
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(1.0))
    tf = txBox.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "Platform Highlights"
    p.font.name = 'Open Sans'
    p.font.color.rgb = RGBColor(192, 0, 0)
    p.font.size = Pt(36)
    p.font.bold = True
    
    rows = 6
    cols = 6
    
    table_left = 0.5
    table_top = 1.2
    
    shape = slide.shapes.add_table(rows, cols, Inches(table_left), Inches(table_top), Inches(12.333), Inches(5.5))
    table = shape.table
    
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(1.5)
    table.columns[2].width = Inches(1.2)
    table.columns[3].width = Inches(1.5)
    table.columns[4].width = Inches(1.2)
    table.columns[5].width = Inches(4.933)
    
    table.rows[0].height = Inches(0.5)
    for i in range(1, 6):
        table.rows[i].height = Inches(1.0)
    
    headers = ["Platform", "Followers", "Posts", "Avg Reach", "Avg ER", "Key Notes/Findings"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(192, 0, 0)
        align = PP_ALIGN.LEFT if i == 5 else PP_ALIGN.CENTER
        _set_cell_text(cell, h, bold=True, color=RGBColor(255, 255, 255), size=Pt(14), align=align)
    
    platforms = ["Facebook", "Instagram", "TikTok", "YouTube", "LinkedIn"]
    
    for idx, platform in enumerate(platforms):
        row_idx = idx + 1
        data = platforms_data.get(platform, {})
        
        bg_color = RGBColor(245, 245, 245) if row_idx % 2 == 1 else RGBColor(255, 255, 255)
        for c in range(6):
            cell = table.cell(row_idx, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_color
            
        p_cell = table.cell(row_idx, 0)
        p_cell.fill.solid()
        p_cell.fill.fore_color.rgb = RGBColor(0, 0, 0)
        _set_cell_text(p_cell, f"        {platform}", bold=True, color=RGBColor(255, 255, 255), size=Pt(14), align=PP_ALIGN.LEFT)
        
        logo_path = download_logo(platform)
        if logo_path:
            logo_top = table_top + 0.5 + (row_idx - 1) * 1.0 + 0.35
            logo_left = table_left + 0.2
            slide.shapes.add_picture(logo_path, Inches(logo_left), Inches(logo_top), height=Inches(0.3))
        
        _set_cell_text(table.cell(row_idx, 1), "N/A", size=Pt(12))
        
        if data:
            _set_cell_text(table.cell(row_idx, 2), f"{data.get('posts', 0):,}", size=Pt(12))
            _add_subtext(table.cell(row_idx, 2), f"({data.get('prev_posts', 0):,} in {prev_month_label})")
            
            _set_cell_text(table.cell(row_idx, 3), f"{data.get('avg_reach', 0):,.0f}", size=Pt(12))
            _add_subtext(table.cell(row_idx, 3), f"({data.get('prev_avg_reach', 0):,.0f} in {prev_month_label})")
            
            _set_cell_text(table.cell(row_idx, 4), f"{data.get('avg_er', 0):.2f}%", size=Pt(12))
            _add_subtext(table.cell(row_idx, 4), f"({data.get('prev_avg_er', 0):.2f}% in {prev_month_label})")
            
            insight = data.get('insight', 'No data available.')
            _set_cell_text(table.cell(row_idx, 5), insight, size=Pt(10), align=PP_ALIGN.LEFT)
        else:
            for c in range(2, 6):
                _set_cell_text(table.cell(row_idx, c), "-", size=Pt(12))

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output
